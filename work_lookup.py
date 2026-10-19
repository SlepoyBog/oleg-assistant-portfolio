import re
from pathlib import Path

from openpyxl import load_workbook


def normalize(value):
    text = str(value or "").casefold().replace("ё", "е")
    text = text.replace("ду", "dn")
    return re.sub(r"[^a-zа-я0-9]+", " ", text).strip()


class WorkLookup:
    """Read-only lookup for approved production reference files."""

    def __init__(self, root):
        self.root = Path(root) if root else None

    def _files(self, relative_dir, pattern):
        directory = self.root / relative_dir if self.root else None
        return sorted(directory.glob(pattern)) if directory and directory.is_dir() else []

    def saw_tables(self, query):
        aliases = {
            "зц": "зц", "цпбк": "цпбк", "цпп": "цпп", "мр": "мр",
        }
        wanted = aliases.get(normalize(query), normalize(query))
        files = self._files(Path("Таблицы распила") / "Актуальные", "*.pdf")
        matches = [path for path in files if wanted in normalize(path.stem)] if wanted else files
        return [self._file_result(path) for path in matches]

    def cutting_sources(self, query):
        text = normalize(query)
        material_aliases = {
            "ст20": "ст20", "ст 20": "ст20",
            "09г2с": "09г2с", "12х18н10т": "12х18н10т",
        }
        material = next((value for key, value in material_aliases.items() if key in text), "")
        operation = "пробивка" if "пробив" in text else "резка" if "резк" in text else ""
        pdfs = self._files(Path("Режимы резки") / "Режимы резки Пдф", "*.pdf")
        matches = []
        for path in pdfs:
            name = normalize(path.stem)
            if material and material not in name:
                continue
            if operation and operation not in name:
                continue
            matches.append(self._file_result(path))
        return matches

    def cutting_rows(self, query, limit=5):
        text = normalize(query)
        dn_match = re.search(r"dn\s*(\d+)\s+(\d+)", text)
        if not dn_match or not self.root:
            return []
        dn_key = f"dn{dn_match.group(1)} {dn_match.group(2)}"
        books = sorted((self.root / "Режимы резки").glob("*.xlsx"))
        if not books:
            return []

        material = next((name for name in ("ст20", "09г2с", "12х18н10т") if name in text), "")
        workbook = load_workbook(books[0], read_only=True, data_only=True)
        results = []
        try:
            for sheet in workbook.worksheets[:3]:
                sheet_material = normalize(sheet.title).replace(" ", "")
                if material and material != sheet_material:
                    continue
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if dn_key not in normalize(row[0]):
                        continue
                    results.append({
                        "part": row[0], "material": row[1], "thickness": row[4],
                        "nozzle": row[5], "speed": row[6], "pressure": row[9],
                        "power": row[10], "focus": row[13], "status": row[24],
                        "source": books[0].name,
                    })
                    if len(results) >= limit:
                        return results
        finally:
            workbook.close()
        return results

    @staticmethod
    def _file_result(path):
        date_match = re.search(r"(\d{2}\.\d{2}\.\d{4})", path.stem)
        return {
            "name": path.name,
            "date": date_match.group(1) if date_match else "дата в имени не указана",
            "status": "утверждено",
        }
