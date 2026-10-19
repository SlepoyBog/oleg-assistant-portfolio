import shutil
from datetime import date, time
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook

from plan_exporter import export_management_plan
from storage import Storage


def create_template(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Мой распорядок"
    sheet.append(["Мой распорядок"])
    sheet["G2"] = ""
    sheet.append(["Начало", "Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"])
    for row, minutes in enumerate(range(8 * 60, 17 * 60 + 1, 15), start=4):
        sheet.cell(row=row, column=1).value = time(minutes // 60, minutes % 60)
        sheet.cell(row=row, column=1).number_format = "h:mm"
    for column in range(2, 7):
        sheet.merge_cells(start_row=22, start_column=column, end_row=25, end_column=column)
        sheet.cell(row=22, column=column).value = "Обед"
    sheet.cell(row=57, column=1).value = "ID"
    workbook.save(path)


root = Path(__file__).resolve().parent / ".tmp_plan_export"
shutil.rmtree(root, ignore_errors=True)
root.mkdir(parents=True)
try:
    template = root / "template.xlsx"
    target = root / "plan.xlsx"
    create_template(template)

    storage = Storage(root / "assistant.db")
    storage.initialize()
    task_id, _ = storage.add_entry("task", "Проверить станок")
    storage.set_task_duration(task_id, 30)

    export_management_plan(storage, template, target, today=date(2026, 9, 30))
    workbook = load_workbook(target)
    sheet = workbook["Мой распорядок"]
    assert sheet["A58"].value == task_id
    assert sheet["C58"].value == "Проверить станок"
    assert sheet["D58"].value.date() == date(2026, 9, 30)
    assert sheet["E58"].value == "08:00"
    assert sheet["F58"].value == 30
    assert sheet["H58"].value == "Нет"
    assert sheet["D4"].value == "Проверить станок"
    workbook.close()

    workbook = load_workbook(target)
    sheet = workbook["Мой распорядок"]
    sheet["D58"] = date(2026, 10, 1)
    sheet["E58"] = "14:15"
    sheet["F58"] = 45
    sheet["H58"] = "Да"
    workbook.save(target)
    workbook.close()

    export_management_plan(storage, template, target, today=date(2026, 9, 30))
    task = storage.entry(task_id)
    assert task["planned_for"] == "2026-10-01"
    assert task["scheduled_at"] == "14:15"
    assert task["duration_minutes"] == 45
    assert task["schedule_locked"] == 1

    updated = load_workbook(target)
    sheet = updated["Мой распорядок"]
    assert sheet["D58"].value.date() == date(2026, 10, 1)
    assert sheet["E58"].value == "14:15"
    assert sheet["H58"].value == "Да"
    assert sheet["E29"].value == "Проверить станок"
    updated.close()

    workbook = load_workbook(target)
    sheet = workbook["Мой распорядок"]
    sheet["H58"] = "Нет"
    workbook.save(target)
    workbook.close()

    export_management_plan(storage, template, target, today=date(2026, 9, 30))
    task = storage.entry(task_id)
    assert task["schedule_locked"] == 0
    assert task["planned_for"] == "2026-10-01"

    workbook = load_workbook(target)
    sheet = workbook["Мой распорядок"]
    sheet["E58"] = None
    sheet["F58"] = None
    workbook.save(target)
    workbook.close()

    export_management_plan(storage, template, target, today=date(2026, 9, 30))

    workbook = load_workbook(target)
    sheet = workbook["Мой распорядок"]
    sheet["A59"] = None
    sheet["C59"] = "Ручная задача без ID"
    workbook.save(target)
    workbook.close()

    try:
        export_management_plan(storage, template, target, today=date(2026, 9, 30))
    except ValueError as error:
        assert "59" in str(error)
        assert "ID" in str(error)
    else:
        raise AssertionError("Строка без ID должна останавливать экспорт")

    archives = list((root / "Архив").glob("*.xlsx"))
    assert len(archives) == 3
finally:
    storage = None
    shutil.rmtree(root, ignore_errors=True)

print("all checks passed")
