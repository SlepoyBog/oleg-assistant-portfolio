import tempfile
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from exporter import export_workbook
from storage import Storage


with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    storage = Storage(root / "assistant.db")
    storage.initialize()
    storage.add_entry("note", "Размер тары 1140х740")
    storage.add_entry("task", "Проверить защитное стекло")
    storage.save_card("part", "Корпус 150_125", {"material": "09Г2С"})
    storage.inventory_change("защитное стекло", 20)
    storage.inventory_change("защитное стекло", -2, "Longxin T28")

    target = export_workbook(storage, root / "oleg-assistant.xlsx")
    workbook = load_workbook(target, data_only=False)

    assert workbook.sheetnames == [
        "Сводка", "Заметки", "Задачи", "Карточки", "Остатки", "История склада"
    ]
    assert workbook["Сводка"]["B3"].value == 1
    assert workbook["Заметки"]["C2"].value == "Размер тары 1140х740"
    assert workbook["Остатки"]["B2"].value == 18
    assert workbook["История склада"]["C2"].value == -2
    assert workbook["Сводка"].freeze_panes == "A3"

print("all checks passed")
