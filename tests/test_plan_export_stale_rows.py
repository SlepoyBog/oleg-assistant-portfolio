import shutil
from datetime import date
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from plan_exporter import export_management_plan
from storage import Storage


project = Path(__file__).resolve().parents[1]
root = Path(__file__).resolve().parent / ".tmp_plan_export_stale"
shutil.rmtree(root, ignore_errors=True)
root.mkdir(parents=True)
try:
    storage = Storage(root / "assistant.db")
    storage.initialize()
    task_id, _ = storage.add_entry("task", "Задача для завершения")
    storage.set_task_duration(task_id, 30)
    event_id, _ = storage.add_entry(
        "event",
        "Совещание",
        planned_for="2026-09-30",
        scheduled_at="10:00",
    )

    target = root / "plan.xlsx"
    template = project / "templates" / "weekly-plan.xlsx"
    export_management_plan(storage, template, target, today=date(2026, 9, 30))

    assert storage.complete(task_id)
    export_management_plan(storage, template, target, today=date(2026, 9, 30))

    workbook = load_workbook(target)
    sheet = workbook["Мой распорядок"]
    ids = [sheet.cell(row=row, column=1).value for row in range(58, 74)]
    assert task_id not in ids
    assert event_id in ids
    workbook.close()

    assert storage.delete_entry(event_id)
    export_management_plan(storage, template, target, today=date(2026, 9, 30))
finally:
    storage = None
    shutil.rmtree(root, ignore_errors=True)

print("all checks passed")
