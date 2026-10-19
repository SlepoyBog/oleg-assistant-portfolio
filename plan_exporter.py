from copy import copy
from datetime import date, datetime, time, timedelta
import os
from pathlib import Path
import shutil
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from scheduler import LUNCH_END, LUNCH_START, Planner, SLOT_MINUTES, WORK_END, WORK_START
from storage import CATEGORY_NAMES


SHEET_NAME = "Мой распорядок"
SCHEDULE_FIRST_ROW = 4
SCHEDULE_LAST_ROW = 40
TASK_HEADER_ROW = 57
TASK_FIRST_ROW = 58
TASK_LAST_ROW = 73
TASK_HEADERS = (
    "ID", "Приоритет", "Задача / проект", "Дата", "Начало",
    "Длительность, мин", "Крайний срок", "Закреплено", "Статус", "Категория", "Готово, %",
)


def export_management_plan(storage, template_path, target, today=None):
    template_path = Path(template_path)
    target = Path(target)
    if not template_path.is_file():
        raise FileNotFoundError(f"Шаблон расписания не найден: {template_path}")

    today = today or date.today()
    week_start = today - timedelta(days=today.weekday())
    week_end = week_start + timedelta(days=6)

    if target.is_file():
        _import_manual_plan(storage, target)

    rows = [dict(row) for row in storage.scheduling_entries()]
    result = Planner().build(rows, week_start, not_before=today)
    storage.apply_schedule(result.allocations)
    rows = [dict(row) for row in storage.report_entries()]

    workbook = load_workbook(template_path)
    sheet = workbook[SHEET_NAME]
    sheet["G2"] = week_start
    sheet["G2"].number_format = "dd.mm.yyyy"

    _clear_schedule(sheet)
    _fill_schedule(sheet, rows, week_start, week_end)
    _fill_task_table(sheet, rows, today, result.unscheduled)

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.stem}-{uuid4().hex}.tmp.xlsx")
    workbook.save(temporary)
    archive = None
    try:
        archive = _archive_existing(target)
        os.replace(temporary, target)
    except Exception:
        temporary.unlink(missing_ok=True)
        if archive:
            archive.unlink(missing_ok=True)
        raise
    return target


def _import_manual_plan(storage, target):
    workbook = load_workbook(target, data_only=False)
    try:
        sheet = workbook[SHEET_NAME]
        legacy_layout = sheet.cell(row=TASK_HEADER_ROW, column=7).value == "Закреплено"
        for row_number in range(TASK_FIRST_ROW, TASK_LAST_ROW + 1):
            entry_id = sheet.cell(row=row_number, column=1).value
            text = sheet.cell(row=row_number, column=3).value
            if not entry_id:
                if text:
                    raise ValueError(
                        f"Ручная строка {row_number} не имеет ID задачи; "
                        "автоматическое обновление остановлено"
                    )
                continue
            if not isinstance(entry_id, int):
                raise ValueError(f"Некорректный ID задачи в строке {row_number}")
            entry = storage.entry(entry_id)
            if (
                entry is None
                or entry["status"] != "open"
                or entry["entry_type"] != "task"
            ):
                continue
            if legacy_layout:
                locked = _is_yes(sheet.cell(row=row_number, column=7).value)
            else:
                deadline = _cell_date(sheet.cell(row=row_number, column=7).value)
                storage.set_task_deadline(
                    entry_id, deadline.isoformat() if deadline else None
                )
                locked = _is_yes(sheet.cell(row=row_number, column=8).value)
            planned = _cell_date(sheet.cell(row=row_number, column=4).value)
            start = _cell_time(sheet.cell(row=row_number, column=5).value)
            duration = sheet.cell(row=row_number, column=6).value
            if not locked and (
                planned is None or start is None or duration in (None, "")
            ):
                continue
            if planned is None or start is None or duration in (None, ""):
                raise ValueError(
                    f"Для задачи #{entry_id} нужны дата, начало и длительность"
                )
            duration = int(duration)
            _validate_manual_slot(start, duration)
            if not storage.set_task_schedule(
                entry_id,
                planned_for=planned.isoformat(),
                scheduled_at=start,
                duration_minutes=duration,
                locked=locked,
            ):
                raise ValueError(f"Открытая задача #{entry_id} из Excel не найдена")
    finally:
        workbook.close()


def _is_yes(value):
    if value is True or value == 1:
        return True
    return str(value or "").strip().casefold() in {"да", "yes", "закреплено"}


def _cell_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        return date.fromisoformat(str(value))
    return None


def _cell_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if value:
        parsed = datetime.strptime(str(value).strip(), "%H:%M")
        return parsed.strftime("%H:%M")
    return None


def _validate_manual_slot(start, duration):
    hours, minutes = map(int, start.split(":"))
    start_minutes = hours * 60 + minutes
    end_minutes = start_minutes + duration
    if duration <= 0 or duration % SLOT_MINUTES:
        raise ValueError("Длительность должна быть положительной и кратной 15 минутам")
    if start_minutes % SLOT_MINUTES or start_minutes < WORK_START or end_minutes > WORK_END:
        raise ValueError("Закреплённая задача выходит за рабочее время 08:00–17:00")
    if start_minutes < LUNCH_END and end_minutes > LUNCH_START:
        raise ValueError("Закреплённая задача пересекается с обедом 12:30–13:30")


def _archive_existing(target, now=None):
    if not target.exists():
        return None
    now = now or datetime.now()
    archive_dir = target.parent / "Архив"
    archive_dir.mkdir(parents=True, exist_ok=True)
    timestamp = now.strftime("%Y-%m-%d %H-%M-%S")
    archive = archive_dir / f"{target.stem} {timestamp}{target.suffix}"
    counter = 2
    while archive.exists():
        archive = archive_dir / f"{target.stem} {timestamp}-{counter}{target.suffix}"
        counter += 1
    shutil.copy2(target, archive)
    return archive


def _clear_schedule(sheet):
    lunch_merges = set()
    for merged in list(sheet.merged_cells.ranges):
        if (
            merged.min_row >= SCHEDULE_FIRST_ROW
            and merged.max_row <= SCHEDULE_LAST_ROW
            and merged.min_col >= 2
            and merged.max_col <= 8
        ):
            top_left = sheet.cell(merged.min_row, merged.min_col).value
            if top_left == "Обед":
                lunch_merges.add(str(merged))
            else:
                sheet.unmerge_cells(str(merged))
    for row in sheet.iter_rows(
        min_row=SCHEDULE_FIRST_ROW,
        max_row=SCHEDULE_LAST_ROW,
        min_col=2,
        max_col=8,
    ):
        for cell in row:
            if isinstance(cell, MergedCell) or cell.value == "Обед":
                continue
            cell.value = None


def _fill_schedule(sheet, rows, week_start, week_end):
    for entry in rows:
        if not entry.get("planned_for") or not entry.get("scheduled_at"):
            continue
        planned = date.fromisoformat(entry["planned_for"])
        if not week_start <= planned <= week_end:
            continue
        hours, minutes = map(int, entry["scheduled_at"].split(":"))
        minute_of_day = hours * 60 + minutes
        if minute_of_day < WORK_START or minute_of_day >= WORK_END:
            continue
        duration = entry.get("duration_minutes") or SLOT_MINUTES
        row_number = SCHEDULE_FIRST_ROW + (minute_of_day - WORK_START) // SLOT_MINUTES
        row_count = max(1, (duration + SLOT_MINUTES - 1) // SLOT_MINUTES)
        end_row = min(SCHEDULE_LAST_ROW, row_number + row_count - 1)
        column_number = 2 + planned.weekday()
        if column_number > 8:
            continue
        start_cell = sheet.cell(row_number, column_number)
        if isinstance(start_cell, MergedCell) or start_cell.value == "Обед":
            continue
        if end_row > row_number:
            sheet.merge_cells(
                start_row=row_number,
                start_column=column_number,
                end_row=end_row,
                end_column=column_number,
            )
        start_cell = sheet.cell(row_number, column_number)
        start_cell.value = entry["text"]
        alignment = copy(start_cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "center"
        alignment.horizontal = "center"
        start_cell.alignment = alignment


def _fill_task_table(sheet, rows, today, unscheduled):
    for column, header in enumerate(TASK_HEADERS, start=1):
        sheet.cell(row=TASK_HEADER_ROW, column=column).value = header
    for row_number in range(TASK_FIRST_ROW, TASK_LAST_ROW + 1):
        for column_number in range(1, len(TASK_HEADERS) + 1):
            sheet.cell(row=row_number, column=column_number).value = None

    capacity = TASK_LAST_ROW - TASK_FIRST_ROW + 1
    for index, entry in enumerate(rows[:capacity], start=TASK_FIRST_ROW):
        planned = date.fromisoformat(entry["planned_for"]) if entry["planned_for"] else None
        priority = entry.get("priority") or "normal"
        priority_label = {"high": "Высокий", "normal": "Средний", "low": "Низкий"}.get(
            priority, "Средний"
        )
        status = "Не запланировано" if entry["id"] in unscheduled else "Запланировано"
        values = (
            entry["id"],
            priority_label,
            entry["text"],
            planned,
            entry.get("scheduled_at"),
            entry.get("duration_minutes"),
            date.fromisoformat(entry["deadline"]) if entry.get("deadline") else None,
            "Да" if entry.get("schedule_locked") else "Нет",
            status,
            CATEGORY_NAMES.get(entry["category"], entry["category"]),
            0,
        )
        for column_number, value in enumerate(values, start=1):
            sheet.cell(row=index, column=column_number).value = value
        sheet.cell(row=index, column=4).number_format = "dd.mm.yyyy"
        sheet.cell(row=index, column=7).number_format = "dd.mm.yyyy"
        sheet.cell(row=index, column=11).number_format = "0%"

    if len(rows) > capacity:
        sheet.cell(row=TASK_LAST_ROW, column=10).value = (
            f"Показаны первые {capacity} из {len(rows)} записей"
        )
