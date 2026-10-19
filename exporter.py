from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")


def export_workbook(storage, target):
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)

    with storage.connect() as connection:
        entries = connection.execute(
            "SELECT id,kind,text,category,status,created_at,completed_at "
            "FROM entries ORDER BY id"
        ).fetchall()
        cards = connection.execute(
            "SELECT id,card_type,name,machine,material,thickness,notes,updated_at "
            "FROM cards ORDER BY name"
        ).fetchall()
        inventory = connection.execute(
            "SELECT id,name,quantity,min_quantity,unit,updated_at "
            "FROM inventory ORDER BY name"
        ).fetchall()
        history = connection.execute(
            "SELECT inventory_events.id,inventory.name,inventory_events.delta,"
            "inventory_events.machine,inventory_events.note,inventory_events.created_at "
            "FROM inventory_events JOIN inventory "
            "ON inventory.id=inventory_events.item_id "
            "ORDER BY inventory_events.id DESC"
        ).fetchall()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    summary.append(["Oleg Assistant — экспорт данных"])
    summary.append(["Дата экспорта", datetime.now().replace(microsecond=0)])
    summary.append(["Заметки", sum(row["kind"] == "note" for row in entries)])
    summary.append([
        "Открытые задачи",
        sum(row["kind"] == "task" and row["status"] == "open" for row in entries),
    ])
    summary.append(["Карточки", len(cards)])
    summary.append(["Расходники", len(inventory)])
    summary.freeze_panes = "A3"
    summary["A1"].font = Font(size=16, bold=True, color="1F1F1F")
    summary["A1"].fill = TITLE_FILL
    summary.merge_cells("A1:B1")
    summary["B2"].number_format = "dd.mm.yyyy hh:mm"

    notes = [row for row in entries if row["kind"] == "note"]
    tasks = [row for row in entries if row["kind"] == "task"]
    _add_sheet(
        workbook,
        "Заметки",
        ["ID", "Категория", "Текст", "Создано"],
        [[row["id"], row["category"], row["text"], row["created_at"]] for row in notes],
    )
    _add_sheet(
        workbook,
        "Задачи",
        ["ID", "Категория", "Текст", "Статус", "Создано", "Выполнено"],
        [[
            row["id"], row["category"], row["text"], row["status"],
            row["created_at"], row["completed_at"],
        ] for row in tasks],
    )
    _add_sheet(
        workbook,
        "Карточки",
        ["ID", "Тип", "Название", "Станок", "Материал", "Толщина", "Примечание", "Обновлено"],
        [[
            row["id"], row["card_type"], row["name"], row["machine"],
            row["material"], row["thickness"], row["notes"], row["updated_at"],
        ] for row in cards],
    )
    _add_sheet(
        workbook,
        "Остатки",
        ["ID", "Количество", "Минимум", "Единица", "Наименование", "Обновлено"],
        [[
            row["id"], row["quantity"], row["min_quantity"], row["unit"],
            row["name"], row["updated_at"],
        ] for row in inventory],
    )
    _add_sheet(
        workbook,
        "История склада",
        ["ID", "Наименование", "Изменение", "Станок", "Примечание", "Дата"],
        [[
            row["id"], row["name"], row["delta"], row["machine"],
            row["note"], row["created_at"],
        ] for row in history],
    )

    _fit_columns(summary)
    workbook.save(target)
    return target


def _add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    _fit_columns(sheet)
    return sheet


def _fit_columns(sheet):
    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[letter].width = min(max(width, 10), 45)
